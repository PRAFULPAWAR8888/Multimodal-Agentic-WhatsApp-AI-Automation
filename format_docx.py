import fitz
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

def main():
    pdf_path = "Agentic_WhatsApp_Platform_Technical_Overview_v4.pdf"
    docx_path = "Letter-Head.docx"
    output_path = "Letter-Head_Updated.docx"

    print("Extracting and formatting PDF content...")
    try:
        pdf_doc = fitz.open(pdf_path)
    except Exception as e:
        print(f"Error opening PDF: {e}")
        return

    try:
        document = Document(docx_path)
    except Exception as e:
        print(f"Error opening DOCX: {e}")
        return

    # Clear body
    for p in document.paragraphs:
        p._element.getparent().remove(p._element)
    for t in document.tables:
        t._element.getparent().remove(t._element)

    for page in pdf_doc:
        blocks = page.get_text("dict")["blocks"]
        for b in blocks:
            if b['type'] == 0:  # text
                text = ""
                is_bold = False
                is_heading = False
                
                # Heuristics for headings: font size > 11.5 or bold
                for line in b["lines"]:
                    for span in line["spans"]:
                        span_text = span["text"]
                        if not span_text.strip():
                            continue
                        text += span_text + " "
                        
                        if span["size"] > 11.5:
                            is_heading = True
                        if "bold" in span["font"].lower():
                            is_bold = True
                            
                text = text.strip()
                # Clean up multiple spaces
                text = re.sub(r'\s+', ' ', text)
                
                if not text:
                    continue
                
                # Detect list items
                is_bullet = text.startswith("•") or text.startswith("-") or text.startswith("") or text.startswith("o ")
                if is_bullet:
                    # Remove bullet character
                    text = text[1:].strip()

                if is_heading or (is_bold and len(text) < 100 and not text.endswith('.')):
                    p = document.add_paragraph()
                    run = p.add_run(text)
                    run.font.name = 'Arial'
                    run.font.size = Pt(14)
                    run.bold = True
                    # Professional Dark Blue color for headings
                    run.font.color.rgb = RGBColor(0, 51, 102) 
                    
                    p.paragraph_format.space_before = Pt(18)
                    p.paragraph_format.space_after = Pt(6)
                    p.paragraph_format.keep_with_next = True
                elif is_bullet:
                    try:
                        p = document.add_paragraph(text, style='List Bullet')
                    except KeyError:
                        # Fallback if style doesn't exist
                        p = document.add_paragraph(f"• {text}")
                        p.paragraph_format.left_indent = Inches(0.25)
                        
                    for run in p.runs:
                        run.font.name = 'Arial'
                        run.font.size = Pt(11)
                    
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.line_spacing = 1.15
                else:
                    p = document.add_paragraph()
                    run = p.add_run(text)
                    run.font.name = 'Arial'
                    run.font.size = Pt(11)
                    
                    p.paragraph_format.space_before = Pt(6)
                    p.paragraph_format.space_after = Pt(12)
                    p.paragraph_format.line_spacing = 1.15
                    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    print("Saving professional DOCX...")
    document.save(output_path)
    print("Format complete!")

if __name__ == "__main__":
    main()
