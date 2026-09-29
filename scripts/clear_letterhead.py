from docx import Document

def delete_paragraph(paragraph):
    p = paragraph._element
    p.getparent().remove(p)
    paragraph._p = paragraph._element = None

def clear_doc():
    filepath = 'Letter-Head.docx'
    doc = Document(filepath)
    
    # Keep the first 3 paragraphs (which contains the original letterhead contact/title info)
    # Delete all other paragraphs that were added
    while len(doc.paragraphs) > 3:
        delete_paragraph(doc.paragraphs[-1])
        
    # Remove all tables that were added
    for table in doc.tables:
        table._element.getparent().remove(table._element)
        
    doc.save(filepath)
    print("Document cleared successfully.")

if __name__ == "__main__":
    clear_doc()
