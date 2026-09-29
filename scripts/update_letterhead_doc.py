from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

def update_letterhead():
    filepath = 'Letter-Head.docx'
    
    # Open the existing document to preserve headers, footers, and images
    doc = Document(filepath)
    
    # --- CUSTOM STYLES ---
    styles = doc.styles
    
    # Check and add custom Title Style if not exists
    if 'CustomTitle' not in styles:
        title_style = styles.add_style('CustomTitle', WD_STYLE_TYPE.PARAGRAPH)
        title_font = title_style.font
        title_font.name = 'Arial'
        title_font.size = Pt(22)
        title_font.color.rgb = RGBColor(0, 51, 102) # Dark Blue
        title_font.bold = True
    else:
        title_style = styles['CustomTitle']
        
    # Check and add custom Code Style
    if 'CodeBlock' not in styles:
        code_style = styles.add_style('CodeBlock', WD_STYLE_TYPE.PARAGRAPH)
        code_font = code_style.font
        code_font.name = 'Courier New'
        code_font.size = Pt(9.5)
        code_font.color.rgb = RGBColor(40, 44, 52)
    else:
        code_style = styles['CodeBlock']
        
    # Set default heading 1 and 2 properties
    if 'Heading 1' not in styles:
        h1_style = styles.add_style('Heading 1', WD_STYLE_TYPE.PARAGRAPH)
    else:
        h1_style = styles['Heading 1']
    h1_font = h1_style.font
    h1_font.name = 'Arial'
    h1_font.size = Pt(16)
    h1_font.color.rgb = RGBColor(0, 102, 204)
        
    if 'Heading 2' not in styles:
        h2_style = styles.add_style('Heading 2', WD_STYLE_TYPE.PARAGRAPH)
    else:
        h2_style = styles['Heading 2']
    h2_font = h2_style.font
    h2_font.name = 'Arial'
    h2_font.size = Pt(14)
    h2_font.color.rgb = RGBColor(64, 64, 64)
    
    # Add some spacing before starting content
    doc.add_paragraph('\n')
    
    # --- DOCUMENT CONTENT ---
    
    # Header Section
    title = doc.add_paragraph('Multimodal Agentic WhatsApp AI Automation Platform', style='CustomTitle')
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph('Comprehensive Technical Project Documentation', style='Heading 2')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph() # Spacer
    
    # 1. Overview
    doc.add_heading('1. Project Overview', level=1)
    doc.add_paragraph('This document outlines the architecture, completed features, and technical specifications of the Multimodal Agentic WhatsApp AI Automation Platform. The platform provides a robust, highly scalable, and AI-driven automation system capable of intelligently routing customer intents, answering FAQs via RAG (Retrieval-Augmented Generation), extracting leads, and processing multimodal inputs (voice and images) directly through WhatsApp.')
    
    # 2. Architecture
    doc.add_heading('2. System Architecture & Tech Stack', level=1)
    doc.add_paragraph('The foundation is built using modern asynchronous Python libraries to ensure maximum throughput and stability.').italic = True
    
    tech_stack = [
        ('Backend Framework', 'FastAPI (Python 3.11+)'),
        ('Database', 'PostgreSQL (with pgvector for AI memory embeddings)'),
        ('Background Tasks', 'Redis & ARQ Worker'),
        ('LLM Independence', 'OpenAI, Ollama (Local), HuggingFace (Failover)'),
        ('Speech-to-Text (STT)', 'faster-whisper (CPU optimized, int8)'),
        ('Text-to-Speech (TTS)', 'Piper TTS (ONNX models generating OGG/Opus audio)'),
        ('Vision & OCR', 'Moondream2 & PaddleOCR')
    ]
    
    for item, desc in tech_stack:
        p = doc.add_paragraph()
        p.add_run('• ').bold = True
        p.add_run(f'{item}: ').bold = True
        p.add_run(desc)
    
    # 3. Core Components
    doc.add_heading('3. Core AI Agent Network', level=1)
    doc.add_paragraph('The platform utilizes a LangGraph-based workflow where a Supervisor Agent routes incoming WhatsApp messages to specialized sub-agents based on 14 distinct intents:')
    
    agents = [
        ('Supervisor Agent', 'Intent classification, confidence gating, and dynamic workflow routing.'),
        ('Knowledge Agent', 'Answers FAQs using RAG (pgvector cosine similarity) from uploaded PDFs and URLs.'),
        ('Lead Agent', 'Extracts structured data from conversations to qualify leads and save to PostgreSQL.'),
        ('Scheduling Agent', 'Handles appointment bookings (connected via MCP external tools).'),
        ('CRM Agent', 'Retrieves and updates customer tickets in external platforms (HubSpot/Frappe).'),
        ('Media Agent', 'Interprets image contents sent by the user.'),
        ('Sales & Support', 'Specialized conversational flows for pricing and troubleshooting.'),
        ('Human Escalation', 'Automatically halts AI responses and flags the chat for live human takeover.')
    ]
    
    for idx, (agent, role) in enumerate(agents, 1):
        p = doc.add_paragraph()
        p.add_run(f'{idx}. {agent}: ').bold = True
        p.add_run(role)
        
    # 4. Code Example
    doc.add_heading('4. Implementation Showcase: Pydantic Settings', level=1)
    doc.add_paragraph('The following code snippet demonstrates how the system safely manages configuration and feature flags using Python Pydantic models (from src/whatsapp_agent/config/settings.py):')
    
    code_snippet = '''class AppSettings(BaseSettings):
    """Core application settings."""
    app_name: str = Field(default="WhatsApp AI Platform")
    llm_provider: LLMProvider = Field(default=LLMProvider.OPENAI)
    
    # --- Token Governance Limits ---
    llm_max_tokens_per_request: int = Field(default=7000)
    llm_enable_cost_estimation: bool = Field(default=True)
    
    # --- Future-Proof Features ---
    enable_outbound_campaigns: bool = Field(default=False)
    enable_live_sentiment_analysis: bool = Field(default=False)'''
    
    doc.add_paragraph(code_snippet, style='CodeBlock')
    
    # 5. Security
    doc.add_heading('5. Security, Governance & Production Hardening', level=1)
    doc.add_paragraph('The application is production-ready with several critical security layers:')
    
    security = [
        'HMAC-SHA256: Cryptographically verifies all incoming Meta WhatsApp webhooks.',
        'Global Rate Limiting: Prevents API abuse using sliding window logic.',
        'Token Budgets: Implements strict input/output limits per session to prevent LLM cost overruns.',
        'RBAC Tool Registry: External tools are registered with risk levels requiring explicit authentication.'
    ]
    
    for sec in security:
        p = doc.add_paragraph()
        p.add_run('• ').bold = True
        p.add_run(sec)
        
    # Finalize
    doc.add_paragraph('\n\n')
    footer = doc.add_paragraph('Generated automatically by Antigravity AI Assistant')
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].italic = True
    footer.runs[0].font.size = Pt(9)
    
    # Save Document
    doc.save(filepath)
    print(f"Documentation successfully appended to: {filepath}")

if __name__ == "__main__":
    update_letterhead()
