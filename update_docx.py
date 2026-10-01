import fitz # PyMuPDF
from docx import Document
from docx.shared import Pt
import sys

def main():
    pdf_path = "Agentic_WhatsApp_Platform_Technical_Overview_v4.pdf"
    docx_path = "Letter-Head.docx"

    print("Extracting from PDF...")
    try:
        pdf_doc = fitz.open(pdf_path)
    except Exception as e:
        print(f"Error opening PDF: {e}")
        sys.exit(1)

    print("Opening DOCX...")
    try:
        document = Document(docx_path)
    except Exception as e:
        print(f"Error opening DOCX: {e}")
        sys.exit(1)

    # Clear existing paragraphs and tables in the document body
    for p in document.paragraphs:
        p._element.getparent().remove(p._element)

    for t in document.tables:
        t._element.getparent().remove(t._element)

    print("Transferring content...")
    for page_num, page in enumerate(pdf_doc):
        # get text as dictionary to extract font sizes and properties
        blocks = page.get_text("dict")["blocks"]
        for b in blocks:
            if b['type'] == 0:  # text block
                # A block can have multiple lines
                p = document.add_paragraph()
                p.paragraph_format.space_after = Pt(8)
                p.paragraph_format.space_before = Pt(0)
                
                for line in b["lines"]:
                    for span in line["spans"]:
                        text = span["text"]
                        if not text.strip():
                            continue
                        run = p.add_run(text)
                        
                        # Apply font settings
                        run.font.name = 'Calibri'
                        
                        # Scale font size slightly if needed, keeping readable size
                        font_size = span["size"]
                        if font_size > 14:
                            run.font.size = Pt(14)
                            run.bold = True
                        elif font_size > 12:
                            run.font.size = Pt(12)
                            run.bold = True
                        else:
                            run.font.size = Pt(11)
                            
                        # Check bold flag from font name heuristically
                        if "bold" in span["font"].lower():
                            run.bold = True
                            
                # If paragraph ended up empty, remove it
                if not p.text.strip():
                    p._element.getparent().remove(p._element)
    
    print("Saving updated DOCX...")
    document.save("Letter-Head_Updated.docx")
    print("Update complete!")

if __name__ == "__main__":
    main()
